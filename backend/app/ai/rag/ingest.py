"""文件上載 → 抽取 → 切段 → 入索引（T3）。

設計（docs/rag-SPEC.md + rag-TODO.md T3）：
  - 支援：PDF（pymupdf）、txt/md、docx、xlsx —— 用 repo 已有嘅 library，唔加新依賴
  - **All-or-nothing**：抽取／embed 失敗唔會留低半生半死嘅文件 row（先做晒嘢，最後一個
    transaction 插入 file row + vectors）
  - **Dedup**：以內容 SHA-256 判斷同一份文件重複上載（唔會靜默重複入庫）
  - **敏感字串**：password / token / API key 之類唔入索引內容
  - 索引重用 `reindex.py` 嘅切段同 `search.py` 嘅批量 embed（單一 source of truth）
"""

from __future__ import annotations

import hashlib
import io
import logging
import re
from uuid import UUID, uuid4

from sqlalchemy import text
from sqlalchemy.ext.asyncio import AsyncSession

from app.ai.rag.reindex import _chunk_text
from app.ai.rag.search import embed_texts

logger = logging.getLogger(__name__)

MAX_BYTES = 15 * 1024 * 1024  # 15 MB — 超過即拒絕（唔會入索引）
ALLOWED_SUFFIXES = (".pdf", ".txt", ".md", ".docx", ".xlsx")

# 敏感字串：唔入索引內容（spec §7.2 / §10.1 —— restricted 資料唔可以直接寫入 vector content）
_SENSITIVE_RE = re.compile(
    r"(?i)\b(password|passwd|pwd|secret|token|api[_-]?key|authorization|bearer)\b\s*[:=]?\s*\S+"
)

SOURCE_MODULE = "file"


class IngestError(Exception):
    """用戶可見嘅上載錯誤（endpoint 會 map 去 400）。"""


def content_hash(data: bytes) -> str:
    return hashlib.sha256(data).hexdigest()


def redact_sensitive(text_in: str) -> str:
    """遮走敏感字串，並保留一個標記（用戶睇到「有嘢被遮」而唔係靜默消失）。"""
    return _SENSITIVE_RE.sub(lambda m: f"{m.group(1)}:[已遮蔽]", text_in)


def extract_text(data: bytes, filename: str) -> str:
    """按副檔名抽取文字。抽唔到 / 唔支援 → IngestError（用戶可見原因）。"""
    name = (filename or "").lower()
    if not name.endswith(ALLOWED_SUFFIXES):
        raise IngestError(f"不支援嘅格式（只支援 {'、'.join(ALLOWED_SUFFIXES)}）")

    try:
        if name.endswith(".pdf"):
            try:
                import pymupdf as fitz  # PyMuPDF >= 1.24
            except ImportError:  # 舊版只 expose 'fitz'
                import fitz  # type: ignore
            with fitz.open(stream=data, filetype="pdf") as doc:
                text_out = "\n".join(page.get_text() for page in doc)
        elif name.endswith(".docx"):
            from docx import Document  # python-docx
            d = Document(io.BytesIO(data))
            text_out = "\n".join(p.text for p in d.paragraphs)
        elif name.endswith(".xlsx"):
            import openpyxl
            wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
            rows: list[str] = []
            for ws in wb.worksheets:
                for row in ws.iter_rows(values_only=True):
                    cells = [str(c) for c in row if c is not None]
                    if cells:
                        rows.append(" | ".join(cells))
            text_out = "\n".join(rows)
        else:  # .txt / .md
            text_out = data.decode("utf-8", errors="replace")
    except IngestError:
        raise
    except Exception as exc:
        raise IngestError(f"檔案讀取失敗：{type(exc).__name__}") from exc

    if len(text_out.strip()) < 20:
        raise IngestError("抽取唔到文字（可能係掃描圖檔 — 需要 OCR，暫時未支援）")
    return text_out


async def ingest_file(
    db: AsyncSession,
    *,
    tenant_id: UUID,
    workspace_id: UUID | None,
    user_id: UUID | None,
    filename: str,
    mime: str | None,
    data: bytes,
) -> dict:
    """上載一份文件入知識庫 → 抽文字 → 遮敏感 → 切段 → embed → 入索引。"""
    if not data:
        raise IngestError("空檔案")
    if len(data) > MAX_BYTES:
        raise IngestError(f"檔案太大（上限 {MAX_BYTES // (1024 * 1024)} MB）")

    await db.execute(
        text("SELECT set_config('app.tenant_id', :t, true)"), {"t": str(tenant_id)}
    )

    # ── Dedup：同一內容唔應該入兩次 ──
    digest = content_hash(data)
    existing = (
        await db.execute(
            text(
                "SELECT id FROM nexus_crm.files WHERE tenant_id = :t AND content_hash = :h LIMIT 1"
            ),
            {"t": tenant_id, "h": digest},
        )
    ).scalar()
    if existing:
        return {"status": "duplicate", "file_id": str(existing), "docs": 0, "chunks": 0}

    text_raw = extract_text(data, filename)
    text_clean = redact_sensitive(text_raw)
    redacted = text_clean != text_raw

    chunks = _chunk_text(text_clean)
    if not chunks:
        raise IngestError("切段之後冇內容")

    # 先 embed（最貴、最易失敗）—— 失敗就唔會有 file row 留低
    vectors = await embed_texts(chunks)
    if not vectors or all(all(v == 0.0 for v in vec) for vec in vectors):
        raise IngestError("建立索引失敗（embedding 服務無回應），請稍後再試")

    file_id = uuid4()
    doc_id = uuid4()
    ws = workspace_id or tenant_id
    await db.execute(
        text("""
            INSERT INTO nexus_crm.files
                (id, tenant_id, storage_key, original_filename, mime_type, file_size,
                 uploaded_by, content_hash)
            VALUES (:id, :t, :key, :name, :mime, :size, :uid, :hash)
        """),
        {
            "id": file_id,
            "t": tenant_id,
            "key": f"rag/{tenant_id}/{file_id}",  # tenant-prefixed（spec §10.2）
            "name": filename[:255],
            "mime": (mime or "application/octet-stream")[:100],
            "size": len(data),
            "uid": user_id,
            "hash": digest,
        },
    )
    await db.execute(
        text("""
            INSERT INTO nexus_ai.vector_documents
                (id, tenant_id, workspace_id, visibility_scope, source_module,
                 source_record_id, created_at)
            VALUES (:id, :t, :ws, 'workspace', :module, :rec, now())
        """),
        {"id": doc_id, "t": tenant_id, "ws": ws, "module": SOURCE_MODULE, "rec": file_id},
    )
    for ctext, vec in zip(chunks, vectors):
        await db.execute(
            text("""
                INSERT INTO nexus_ai.vector_document_chunks
                    (id, document_id, chunk_text, embedding, tenant_id, workspace_id, visibility_scope)
                VALUES (:id, :doc, :txt, CAST(:vec AS vector), :t, :ws, 'workspace')
            """),
            {
                "id": uuid4(),
                "doc": doc_id,
                "txt": ctext,
                "vec": str(vec),
                "t": tenant_id,
                "ws": ws,
            },
        )
    await db.commit()

    return {
        "status": "indexed",
        "file_id": str(file_id),
        "docs": 1,
        "chunks": len(chunks),
        "redacted": redacted,
        "chars": len(text_clean),
    }


VALID_CLASSIFICATIONS = ("internal", "confidential", "restricted", "archived")


async def set_file_controls(
    db: AsyncSession,
    *,
    tenant_id: UUID,
    file_id: UUID,
    ai_enabled: bool | None = None,
    classification: str | None = None,
) -> int:
    """T5：切換某份文件嘅 AI 可用性／機密級別（即時影響檢索 —— spec §12.1 第 4 項）。"""
    await db.execute(
        text("SELECT set_config('app.tenant_id', :t, true)"), {"t": str(tenant_id)}
    )
    sets: list[str] = []
    params: dict = {"t": tenant_id, "f": file_id}
    if ai_enabled is not None:
        sets.append("ai_enabled = :ae")
        params["ae"] = bool(ai_enabled)
    if classification is not None:
        if classification not in VALID_CLASSIFICATIONS:
            raise IngestError(f"機密級別只可以係：{', '.join(VALID_CLASSIFICATIONS)}")
        sets.append("classification = :cl")
        params["cl"] = classification
    if not sets:
        return 0

    res = await db.execute(
        text(
            f"""
            UPDATE nexus_ai.vector_documents SET {', '.join(sets)}
            WHERE tenant_id = :t AND source_module = 'file' AND source_record_id = :f
            """
        ),
        params,
    )
    await db.commit()
    return int(getattr(res, "rowcount", 0) or 0)


async def get_citation(
    db: AsyncSession, *, tenant_id: UUID, document_id: UUID
) -> dict | None:
    """T4：引用再驗證（spec §7.2）—— 文件仍然存在、未刪除、可檢索、屬當前 tenant。

    取唔到 → None（endpoint 回 404，唔會透露其他 tenant 嘅文件存在與否）。
    """
    await db.execute(
        text("SELECT set_config('app.tenant_id', :t, true)"), {"t": str(tenant_id)}
    )
    row = (
        await db.execute(
            text("""
                SELECT d.id, d.source_module, d.source_record_id, d.ai_enabled,
                       d.document_status, d.classification, f.original_filename
                FROM nexus_ai.vector_documents d
                LEFT JOIN nexus_crm.files f
                       ON d.source_module = 'file' AND f.id = d.source_record_id
                WHERE d.id = :doc AND d.tenant_id = :t
            """),
            {"doc": document_id, "t": tenant_id},
        )
    ).mappings().first()
    if not row:
        return None
    retrievable = (
        bool(row["ai_enabled"])
        and row["document_status"] == "approved"
        and (row["classification"] or "internal") not in ("restricted", "archived", "deleted")
    )
    return {
        "document_id": str(row["id"]),
        "source_module": row["source_module"],
        "source_record_id": str(row["source_record_id"]),
        "title": row["original_filename"] or row["source_module"],
        "retrievable": retrievable,
    }


async def list_files(db: AsyncSession, *, tenant_id: UUID) -> list[dict]:
    await db.execute(
        text("SELECT set_config('app.tenant_id', :t, true)"), {"t": str(tenant_id)}
    )
    rows = (
        await db.execute(
            text("""
                SELECT f.id, f.original_filename, f.mime_type, f.file_size, f.created_at,
                       (SELECT count(*) FROM nexus_ai.vector_document_chunks c
                          JOIN nexus_ai.vector_documents d ON d.id = c.document_id
                         WHERE d.source_module = 'file' AND d.source_record_id = f.id) AS chunks
                FROM nexus_crm.files f
                WHERE f.tenant_id = :t AND f.content_hash IS NOT NULL
                ORDER BY f.created_at DESC
            """),
            {"t": tenant_id},
        )
    ).mappings().all()
    return [dict(r) for r in rows]


async def delete_file(db: AsyncSession, *, tenant_id: UUID, file_id: UUID) -> int:
    """刪文件 = 即刻停止檢索（vector doc 走，chunks 靠 FK CASCADE）。"""
    await db.execute(
        text("SELECT set_config('app.tenant_id', :t, true)"), {"t": str(tenant_id)}
    )
    res = await db.execute(
        text("""
            DELETE FROM nexus_ai.vector_documents
            WHERE tenant_id = :t AND source_module = 'file' AND source_record_id = :f
        """),
        {"t": tenant_id, "f": file_id},
    )
    await db.execute(
        text("DELETE FROM nexus_crm.files WHERE tenant_id = :t AND id = :f"),
        {"t": tenant_id, "f": file_id},
    )
    await db.commit()
    return int(getattr(res, "rowcount", 0) or 0)
